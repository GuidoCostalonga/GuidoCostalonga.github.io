"""Estrae le serie degli archivi parziali FAOSTAT e OCSE presenti nel repository.

Due impaginazioni di archivio:
  lotti 5, 9, 11, 15, 18   -> L<n>/PAESI/<ISO3>/*.xlsx e L<n>/AGGREGATI/<cod>/*.xlsx
  lotti 19 e successivi    -> L<n>/L<n>_<ISO3 o codice>_<nome>.xlsx
Ogni foglio DATI_* ha le 16 colonne comuni (B:Q) e, da R, i campi originali della fonte.
Scrive un file JSON per lotto: <uscita>_L<n>.json

Uso: python3 estrai_archivi_parziali.py <cartella con L5/ L9/ ...> <prefisso uscita> [aree_faostat.json]
"""
import glob, json, os, re, sys
import openpyxl

radice, uscita = sys.argv[1], sys.argv[2]
mappa = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else {}
SALTA_LOTTI = {25}  # FDIQ: schema per indagine e livello geografico, senza annualità
SALTA = re.compile(r'(code|CODE|Code \(|_ID$|^STRUCTURE|^ACTION$|^Year|^Note$|^Unit$|^Area$|Reporter Countries|^Recipient Country$)', re.I)

per_lotto = {}
for d in sorted(glob.glob(radice + '/L*/')):
    per_lotto[int(os.path.basename(d.rstrip('/'))[1:])] = sorted(glob.glob(d + '**/*.xlsx', recursive=True))

for lotto, file in sorted(per_lotto.items()):
    if lotto in SALTA_LOTTI:
        print('lotto', lotto, 'saltato', flush=True)
        continue
    out = {}
    for f in file:
        base = os.path.basename(f)[:-5]
        if not re.match(r'^L\d+_', base):
            continue
        parti = os.path.relpath(f, radice).split('/')
        chiave, agg = None, False
        if len(parti) >= 4 and parti[1] in ('PAESI', 'AGGREGATI'):
            chiave, agg = parti[2], parti[1] == 'AGGREGATI'
        else:
            tok = base.split('_')
            agg = tok[1] in ('AGG', 'DEST') or 'excluding' in base or 'unspecified' in base
            # il token è un codice ISO3 solo per i file dei paesi: gli aggregati (AGG, DEST) prendono la chiave dal codice FAOSTAT
            if re.match(r'^[A-Z]{3}$', tok[1]) and tok[1] not in ('AGG', 'DEST'):
                chiave = tok[1]
        wb = openpyxl.load_workbook(f, read_only=True)
        for s in wb.sheetnames:
            if not s.startswith('DATI_'):
                continue
            righe = wb[s].iter_rows(min_row=3, values_only=True)
            testa = next(righe)[1:]
            extra = [(i, h) for i, h in enumerate(testa) if i >= 16 and h and not SALTA.search(str(h))]
            for r in righe:
                r = r[1:]
                if not r or r[0] is None:
                    continue
                codgeo, nomegeo, ds, ds_n, anno, val, _, unita, _, _, _, stato, agg_f, url, _, ident = r[:16]
                if anno is None:
                    continue
                k = chiave or mappa.get(str(codgeo)) or ('FAO_' + str(codgeo))
                d = out.setdefault(ds, {'n': ds_n, 'url': url, 'agg': agg_f, 'serie': {}, 'terr': {}, 'tipo': {}})
                d['terr'].setdefault(k, nomegeo)
                d['tipo'].setdefault(k, 'agg' if agg else 'paese')
                se = d['serie'].get(ident)
                if se is None:
                    etich = ' · '.join(str(r[i]) for i, h in extra if r[i] not in (None, '', 'Total', 'Not applicable', '_Z', '_T'))
                    se = d['serie'][ident] = {'l': etich, 'u': unita, 'd': {}}
                se['d'].setdefault(k, {})[str(anno)] = [val, stato]
        wb.close()
        print(f, flush=True)
    json.dump({f'{lotto}|{ds}': d for ds, d in out.items()}, open(f'{uscita}_L{lotto}.json', 'w'))
    print('lotto', lotto, 'scritto:', len(out), 'dataset', flush=True)
print('fine')
